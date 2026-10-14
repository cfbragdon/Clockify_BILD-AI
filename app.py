"""Clockify BILD-AI: a small personal time tracker.

Start it with "Start Clockify.bat". Data is stored in clockify.db next to this file.
Times are stored in UTC and shown in local time.
"""
import io
import logging
import socket
import sqlite3
import threading
import webbrowser
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path

from flask import Flask, abort, g, jsonify, make_response, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Font

HOST = "127.0.0.1"  # only reachable from this computer
PORT = 5878
URL = f"http://{HOST}:{PORT}/"
DB_PATH = Path(__file__).with_name("clockify.db")

# Colours given to new projects, in order
COLORS = ["#03a9f4", "#e91e63", "#8bc34a", "#ff9800", "#9c27b0",
          "#009688", "#f44336", "#3f51b5", "#795548", "#607d8b"]

SCHEMA = """
CREATE TABLE IF NOT EXISTS projects (
    id    INTEGER PRIMARY KEY,
    name  TEXT NOT NULL UNIQUE COLLATE NOCASE,
    color TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS entries (
    id          INTEGER PRIMARY KEY,
    description TEXT NOT NULL DEFAULT '',
    project_id  INTEGER REFERENCES projects(id),
    start_at    TEXT NOT NULL,  -- UTC, ISO 8601
    end_at      TEXT            -- UTC, ISO 8601; NULL while the timer is running
);
"""

ENTRY_SELECT = """
SELECT e.id, e.description, e.start_at, e.end_at, p.name AS project, p.color
FROM entries e LEFT JOIN projects p ON p.id = e.project_id
"""

app = Flask(__name__)


# ---------- Database helpers ----------

def init_db():
    con = sqlite3.connect(DB_PATH)
    try:
        con.executescript(SCHEMA)
    finally:
        con.close()


def db():
    if "db" not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys = ON")
    return g.db


@app.teardown_appcontext
def close_db(_exc):
    con = g.pop("db", None)
    if con is not None:
        con.close()


def fail(message, status=400):
    abort(make_response(jsonify(error=message), status))


def now_utc():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def to_utc(value):
    """Turn an ISO timestamp from the browser into our stored UTC format."""
    dt = datetime.fromisoformat(value)
    if dt.tzinfo is None:
        raise ValueError("timestamp has no timezone")
    return dt.astimezone(timezone.utc).isoformat(timespec="seconds")


def project_id_for(name):
    """Find a project by name, creating it if it doesn't exist yet."""
    name = (name or "").strip()
    if not name:
        return None
    row = db().execute("SELECT id FROM projects WHERE name = ?", (name,)).fetchone()
    if row:
        return row["id"]
    count = db().execute("SELECT COUNT(*) FROM projects").fetchone()[0]
    cur = db().execute("INSERT INTO projects (name, color) VALUES (?, ?)",
                       (name, COLORS[count % len(COLORS)]))
    return cur.lastrowid


def get_entry(entry_id):
    row = db().execute(ENTRY_SELECT + " WHERE e.id = ?", (entry_id,)).fetchone()
    return dict(row) if row else None


def running_entry():
    row = db().execute(ENTRY_SELECT + " WHERE e.end_at IS NULL").fetchone()
    return dict(row) if row else None


def read_entry_fields(data):
    """Validate a manually entered or edited entry."""
    try:
        start = to_utc(data["start"])
        end = to_utc(data["end"])
    except (KeyError, TypeError, ValueError):
        fail("Please enter a valid date, start time and end time.")
    if end <= start:
        fail("The end time must be after the start time.")
    description = (data.get("description") or "").strip()
    return description, project_id_for(data.get("project")), start, end


# ---------- Pages ----------

@app.get("/")
def index():
    return render_template("index.html")


# ---------- Projects and entries ----------

@app.get("/api/projects")
def list_projects():
    rows = db().execute("SELECT id, name, color FROM projects ORDER BY name").fetchall()
    return jsonify([dict(r) for r in rows])


@app.get("/api/entries")
def list_entries():
    rows = db().execute(ENTRY_SELECT + " WHERE e.end_at IS NOT NULL ORDER BY e.start_at DESC").fetchall()
    return jsonify([dict(r) for r in rows])


@app.post("/api/entries")
def add_entry():
    description, project_id, start, end = read_entry_fields(request.get_json(silent=True) or {})
    cur = db().execute(
        "INSERT INTO entries (description, project_id, start_at, end_at) VALUES (?, ?, ?, ?)",
        (description, project_id, start, end))
    db().commit()
    return jsonify(get_entry(cur.lastrowid)), 201


@app.put("/api/entries/<int:entry_id>")
def edit_entry(entry_id):
    entry = get_entry(entry_id)
    if entry is None or entry["end_at"] is None:
        fail("That entry no longer exists.", 404)
    description, project_id, start, end = read_entry_fields(request.get_json(silent=True) or {})
    db().execute(
        "UPDATE entries SET description = ?, project_id = ?, start_at = ?, end_at = ? WHERE id = ?",
        (description, project_id, start, end, entry_id))
    db().commit()
    return jsonify(get_entry(entry_id))


@app.delete("/api/entries/<int:entry_id>")
def delete_entry(entry_id):
    db().execute("DELETE FROM entries WHERE id = ?", (entry_id,))
    db().commit()
    return jsonify(ok=True)


# ---------- Timer ----------

@app.get("/api/timer")
def get_timer():
    return jsonify(running_entry())


@app.post("/api/timer/start")
def start_timer():
    if running_entry():
        fail("A timer is already running.", 409)
    data = request.get_json(silent=True) or {}
    cur = db().execute(
        "INSERT INTO entries (description, project_id, start_at) VALUES (?, ?, ?)",
        ((data.get("description") or "").strip(), project_id_for(data.get("project")), now_utc()))
    db().commit()
    return jsonify(get_entry(cur.lastrowid)), 201


@app.post("/api/timer/stop")
def stop_timer():
    entry = running_entry()
    if entry is None:
        fail("No timer is running.", 409)
    data = request.get_json(silent=True) or {}
    db().execute(
        "UPDATE entries SET description = ?, project_id = ?, end_at = ? WHERE id = ?",
        ((data.get("description") or "").strip(), project_id_for(data.get("project")),
         now_utc(), entry["id"]))
    db().commit()
    return jsonify(get_entry(entry["id"]))


# ---------- Spreadsheet download ----------

@app.get("/export.xlsx")
def export_xlsx():
    try:
        first = date.fromisoformat(request.args["from"])
        last = date.fromisoformat(request.args["to"])
    except (KeyError, ValueError):
        fail("Please choose a valid date range.")
    if last < first:
        fail("The 'to' date must be on or after the 'from' date.")

    # Date range in local time, converted to UTC for the query
    range_start = datetime.combine(first, time.min).astimezone().astimezone(timezone.utc)
    range_end = datetime.combine(last + timedelta(days=1), time.min).astimezone().astimezone(timezone.utc)
    rows = db().execute(
        ENTRY_SELECT + " WHERE e.end_at IS NOT NULL AND e.start_at >= ? AND e.start_at < ?"
        " ORDER BY e.start_at",
        (range_start.isoformat(timespec="seconds"), range_end.isoformat(timespec="seconds"))).fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = "Time entries"
    ws.append(["Date", "Start", "End", "Duration", "Hours", "Project", "Description"])
    for cell in ws[1]:
        cell.font = Font(bold=True)

    total_hours = 0.0
    for r in rows:
        start = datetime.fromisoformat(r["start_at"]).astimezone()
        end = datetime.fromisoformat(r["end_at"]).astimezone()
        seconds = int((end - start).total_seconds())
        hours = round(seconds / 3600, 2)
        total_hours += seconds / 3600
        ws.append([start.strftime("%Y-%m-%d"), start.strftime("%H:%M"), end.strftime("%H:%M"),
                   f"{seconds // 3600}:{seconds % 3600 // 60:02d}", hours,
                   r["project"] or "", r["description"]])

    ws.append([])
    ws.append(["Total", "", "", "", round(total_hours, 2)])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.cell(ws.max_row, 5).font = Font(bold=True)
    for col, width in zip("ABCDEFG", [12, 8, 8, 10, 8, 24, 50]):
        ws.column_dimensions[col].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer, as_attachment=True,
        download_name=f"clockify_{first.isoformat()}_to_{last.isoformat()}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ---------- Start-up ----------

def already_running():
    try:
        with socket.create_connection((HOST, PORT), timeout=0.5):
            return True
    except OSError:
        return False


def main():
    if already_running():
        print("Clockify BILD-AI is already running. Opening it in your browser.")
        webbrowser.open(URL)
        return

    from werkzeug.serving import make_server

    init_db()
    logging.getLogger("werkzeug").setLevel(logging.ERROR)  # hide per-request logs
    server = make_server(HOST, PORT, app, threaded=True)
    print("Clockify BILD-AI is running at", URL)
    print("Keep this window open while you use the app. Close it to stop the app.")
    threading.Timer(1, webbrowser.open, [URL]).start()
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass


if __name__ == "__main__":
    main()
